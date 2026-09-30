import glob
import math
import os
import re
import sys
from datetime import datetime

sys.stdout.reconfigure(encoding="utf-8")

from dotenv import load_dotenv
from anthropic import Anthropic, APIError, APIStatusError
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

REPORT_PATTERN = "주간보고_*.txt"
MODEL = "claude-sonnet-5-5"

SYSTEM_PROMPT = """당신은 여러 부서의 보고서를 읽고 부서 간 공통 문제를 찾아 해결 방안을 제안하는 기획 담당자입니다.
결과는 반드시 submit_proposal 도구로 제출합니다. 아래 규칙을 지키세요.

[작성 규칙]
- 문제마다 근거가 되는 보고서 이름(출처 파일명 그대로)과 해당 보고서의 내용(원문 그대로 짧게 인용)을 적습니다.
- 둘 이상의 부서 보고서에서 함께 나타나는 문제를 '공통 문제'로 다룹니다. 한 부서에서만 나온 사안은 공통 문제로 세지 않습니다.
- 원문에 없는 해결 방안은 새로운 내용이므로 tag를 "제안"으로 표시합니다. 원문에 이미 있는 조치는 tag를 "원문"으로 하고 source에 출처 파일명을 적습니다.
- 원문에 없는 효과 수치, 예산, 담당자, 확정 일정은 절대 만들지 않습니다.
  기대 효과는 수치 없이 정성적으로만 쓰고, 수치는 원문에 있는 값만 인용합니다.
- 근거가 부족하거나 원문에서 확인할 수 없는 내용(담당자, 예산, 일정, 원인, 승인 여부 등)은 "확인 필요"라고 표시합니다.
- 원문에 있는 일정은 출처를 밝히고 인용할 수 있지만, 새로운 일정을 확정하지 않습니다.

[원문 충실도 규칙]
- 원문이 추측·조건으로 쓴 표현("가능성 있음", "예상", "우려", "확정 필요", "목표")은 그대로 유지합니다.
  "지연되고 있다", "확정되었다"처럼 더 단정적으로 바꾸지 않습니다. 보고서마다 표현이 다르면(예: 영업팀 "가능성", 품질팀 "불가피/예상") 차이를 그대로 밝힙니다.
- 일정 표현("목표일", "확정 필요", "예정" 등)의 출처는 그 표현이 실제로 적힌 보고서로 씁니다. 같은 일정이라도 보고서마다 확정 수준이 다르면 각각 구분해서 적습니다.
- 여러 사안을 날짜로 묶어 요약("모두 ~주에 몰려 있다")하지 않습니다. 사안마다 원문에 날짜가 있는 경우에만 각각 날짜를 적고, 날짜가 없는 사안은 "일정 확인 필요"로 둡니다.
- 생산계획·실적이 문제와 관련되면 같은 보고서의 실적 수치(예: 계획 대비 달성률)도 함께 인용합니다. 단, 실적과 계획 변경 또는 불량 사이의 인과관계는 원문에 없으면 "확인 필요"로 둡니다.
- 부서 수를 셀 때는 각 보고서의 '여러 팀 협업 필요 사항' 항목에 적힌 부서를 모두 포함하고, 어떤 협업 항목 기준인지 밝힙니다.
  (예: 반입 일정·동선 협의는 구매·생산관리·총무·영업, 설비 사양 상호 확인은 구매·품질.) 어느 부서를 단순 "담당"으로 축소해 쓰지 않습니다.

[분량] 슬라이드에 들어갈 내용이므로 짧고 간결하게 씁니다.
- 공통 문제는 핵심 3~5개. 문제 하나당 근거는 2~4개, 인용문은 60자 이내.
- 해결 방안 문장은 각 90자 이내. 현재 상황은 3~4줄.
"""

PROPOSAL_TOOL = {
    "name": "submit_proposal",
    "description": "부서 공통 문제 해결 제안서를 제출한다.",
    "input_schema": {
        "type": "object",
        "properties": {
            "situation": {
                "type": "array",
                "items": {"type": "string"},
                "description": "1. 현재 상황 (3~4줄)",
            },
            "problems": {
                "type": "array",
                "description": "2. 확인된 문제",
                "items": {
                    "type": "object",
                    "properties": {
                        "title": {"type": "string"},
                        "departments": {"type": "array", "items": {"type": "string"}},
                        "evidence": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "properties": {
                                    "report": {"type": "string", "description": "출처 파일명"},
                                    "quote": {"type": "string", "description": "보고서 내용 인용"},
                                },
                                "required": ["report", "quote"],
                            },
                        },
                        "note": {"type": "string", "description": "원인 불명 등 '확인 필요' 사항. 없으면 빈 문자열"},
                    },
                    "required": ["title", "departments", "evidence", "note"],
                },
            },
            "solutions": {
                "type": "array",
                "description": "3. 해결 방안 (problems와 같은 순서, 문제당 하나)",
                "items": {
                    "type": "object",
                    "properties": {
                        "problem_title": {"type": "string"},
                        "items": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "properties": {
                                    "tag": {"type": "string", "enum": ["제안", "원문"]},
                                    "text": {"type": "string"},
                                    "source": {"type": "string", "description": "tag가 원문일 때 출처 파일명, 아니면 빈 문자열"},
                                },
                                "required": ["tag", "text", "source"],
                            },
                        },
                    },
                    "required": ["problem_title", "items"],
                },
            },
            "effects": {
                "type": "array",
                "items": {"type": "string"},
                "description": "4. 기대 효과 (수치 없는 정성적 서술)",
            },
            "known_schedule": {
                "type": "array",
                "description": "5. 원문에 있는 일정만",
                "items": {
                    "type": "object",
                    "properties": {
                        "text": {"type": "string"},
                        "source": {"type": "string"},
                    },
                    "required": ["text", "source"],
                },
            },
            "to_confirm": {
                "type": "array",
                "items": {"type": "string"},
                "description": "5. 확인 필요 사항 (담당자, 예산, 신규 일정 등 원문에 없는 것)",
            },
        },
        "required": ["situation", "problems", "solutions", "effects", "known_schedule", "to_confirm"],
    },
}


# ---------------------------------------------------------------- 입력 / API

def load_api_key() -> str:
    load_dotenv()
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        raise RuntimeError(".env 파일에 ANTHROPIC_API_KEY가 설정되어 있지 않습니다.")
    return api_key


def load_reports(folder: str):
    paths = sorted(glob.glob(os.path.join(folder, REPORT_PATTERN)))
    if not paths:
        raise RuntimeError(
            f"읽을 자료가 없습니다. ('{folder}' 폴더에서 '{REPORT_PATTERN}' 형식의 보고서 파일을 찾지 못했습니다.)"
        )
    reports = []
    for path in paths:
        with open(path, "r", encoding="utf-8") as f:
            content = f.read().strip()
        if not content:
            raise RuntimeError(f"'{os.path.basename(path)}' 파일 내용이 비어 있습니다.")
        reports.append({"filename": os.path.basename(path), "content": content})
    if len(reports) < 2:
        raise RuntimeError("공통 문제를 찾으려면 보고서가 2개 이상 필요합니다.")
    return reports


def build_user_prompt(reports) -> str:
    parts = ["다음은 각 부서의 보고서 원문입니다. 부서 간 공통 문제를 찾아 제안서를 작성해주세요.\n"]
    for r in reports:
        parts.append(f"----- 출처 파일: {r['filename']} -----\n{r['content']}\n")
    return "\n".join(parts)


def generate_proposal(api_key: str, reports) -> dict:
    client = Anthropic(api_key=api_key)
    response = client.messages.create(
        model=MODEL,
        max_tokens=16000,
        system=SYSTEM_PROMPT,
        tools=[PROPOSAL_TOOL],
        messages=[{"role": "user", "content": build_user_prompt(reports)}],
    )
    if response.stop_reason == "max_tokens":
        raise RuntimeError("응답이 길이 제한에 걸려 중간에 잘렸습니다. 보고서 수를 줄이거나 다시 실행해주세요.")
    for block in response.content:
        if getattr(block, "type", None) == "tool_use" and block.name == "submit_proposal":
            return block.input
    raise RuntimeError("Claude로부터 제안서 내용을 받지 못했습니다.")


def check_proposal(data: dict, reports):
    """사람이 확인해야 할 경고 목록을 돌려준다."""
    warnings = []
    known = {r["filename"] for r in reports}
    cited = set()

    if not data.get("problems"):
        warnings.append("확인된 문제가 하나도 없습니다.")

    for p in data.get("problems", []):
        if not p["evidence"]:
            warnings.append(f"문제 '{p['title']}'에 근거가 없습니다.")
        for e in p["evidence"]:
            cited.add(e["report"])
            if e["report"] not in known:
                warnings.append(f"존재하지 않는 보고서 이름이 인용되었습니다: {e['report']}")
            elif e["quote"].strip('"“” ') and _norm(e["quote"]) not in _norm(
                next(r["content"] for r in reports if r["filename"] == e["report"])
            ):
                warnings.append(f"원문과 글자가 다른 인용: [{e['report']}] {e['quote'][:40]}…")

    for r in reports:
        if r["filename"] not in cited:
            warnings.append(f"'{r['filename']}'은(는) 어떤 문제의 근거로도 인용되지 않았습니다.")

    titles = [p["title"] for p in data.get("problems", [])]
    sol_titles = [s["problem_title"] for s in data.get("solutions", [])]
    if titles != sol_titles:
        warnings.append("해결 방안이 확인된 문제와 1:1로 대응하지 않습니다.")

    tags = [i["tag"] for s in data.get("solutions", []) for i in s["items"]]
    if "제안" not in tags:
        warnings.append("'제안' 표시가 하나도 없습니다.")
    for s in data.get("solutions", []):
        for i in s["items"]:
            if i["tag"] == "원문" and i["source"] not in known:
                warnings.append(f"[원문] 표시 항목의 출처가 올바르지 않습니다: {i['text'][:30]}…")

    contents = {r["filename"]: r["content"] for r in reports}
    schedule_like = [(k["text"], k["source"]) for k in data.get("known_schedule", [])]
    schedule_like += [(i["text"], i["source"]) for s in data.get("solutions", [])
                      for i in s["items"] if i["tag"] == "원문"]
    for text, source in schedule_like:
        if source not in known:
            warnings.append(f"출처가 올바르지 않습니다: {text[:30]}…")
            continue
        # 확정 수준을 나타내는 표현은 그 출처 보고서에 실제로 있어야 한다
        for word in ("목표일", "확정", "예정", "가능성", "예상"):
            if word in text and word not in contents[source]:
                warnings.append(f"'{word}' 표현이 출처 {source}에 없습니다(다른 보고서 표현일 수 있음): {text[:40]}…")

    return warnings


def _norm(s: str) -> str:
    return re.sub(r"[\s\"'“”‘’…]+", "", s)


# ---------------------------------------------------------------- PPT 만들기

NAVY = RGBColor(0x1B, 0x26, 0x3B)
INK = RGBColor(0x22, 0x2B, 0x38)
MUTED = RGBColor(0x6B, 0x75, 0x85)
CARD = RGBColor(0xF1, 0xF4, 0xF8)
AMBER = RGBColor(0xE0, 0x8A, 0x1E)
AMBER_TINT = RGBColor(0xFD, 0xF1, 0xDE)
GRAY_TAG = RGBColor(0x8A, 0x94, 0xA3)
TEAL = RGBColor(0x1F, 0x7A, 0x8C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "맑은 고딕"

SW, SH = 13.333, 7.5
MX = 0.6                      # 좌우 여백
BODY_TOP, BODY_BOTTOM = 1.75, 7.0
BODY_W = SW - 2 * MX


def lines_needed(text: str, width_in: float, pt: float) -> int:
    """한글 기준 글자 폭(≈pt)으로 줄 수를 넉넉히 추정한다."""
    chars_per_line = max(1, int(width_in * 72 / (pt * 1.05)))
    return sum(max(1, math.ceil(len(seg) / chars_per_line)) for seg in text.split("\n"))


def text_height(text: str, width_in: float, pt: float) -> float:
    return lines_needed(text, width_in, pt) * pt * 1.35 / 72


def add_text(slide, x, y, w, h, text, pt, color=INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(pt)
    r.font.bold = bold
    r.font.name = FONT
    r.font.color.rgb = color
    return tb


def add_box(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    return s


def set_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def new_slide(prs, section: str, title: str, page_note: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_text(slide, MX, 0.5, 2.0, 0.3, section, 13, TEAL, bold=True)
    add_text(slide, MX, 0.85, BODY_W - 1.5, 0.7, title, 28, NAVY, bold=True)
    if page_note:
        add_text(slide, SW - MX - 1.5, 0.95, 1.5, 0.4, page_note, 13, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def add_footer(slide, text):
    add_text(slide, MX, 7.08, BODY_W, 0.25, text, 10, MUTED)


def flow_pages(items, avail, gap):
    """(높이, 항목) 목록을 화면 높이에 맞춰 여러 쪽으로 나눈다."""
    pages, cur, used = [], [], 0.0
    for h, it in items:
        if cur and used + h > avail:
            pages.append(cur)
            cur, used = [], 0.0
        cur.append((h, it))
        used += h + gap
    if cur:
        pages.append(cur)
    return pages


def build_title_slide(prs, reports, today):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, NAVY)
    add_text(s, 0.9, 2.2, 11.5, 1.0, "부서 공통 문제 해결 제안서", 44, WHITE, bold=True)
    add_text(s, 0.9, 3.35, 11.5, 0.5, "여러 부서 보고서에서 찾은 공통 문제와 해결 방안", 20,
             RGBColor(0xC9, 0xD3, 0xE3))
    names = ", ".join(r["filename"].replace("주간보고_", "").replace(".txt", "") for r in reports)
    add_text(s, 0.9, 5.4, 11.5, 0.9,
             f"작성일 {today}\n분석 대상: {names} ({len(reports)}개 부서 보고서)", 14,
             RGBColor(0xC9, 0xD3, 0xE3))
    add_text(s, 0.9, 6.7, 11.5, 0.4,
             "AI가 보고서 원문만 바탕으로 작성한 초안입니다. 검토 후 사용하세요.", 12,
             RGBColor(0x9A, 0xA7, 0xBD))


def build_situation(prs, data, reports):
    s = new_slide(prs, "1", "현재 상황")
    # 왼쪽: 상황 요약
    left_w = 7.6
    add_box(s, MX, BODY_TOP, left_w, 4.9, CARD)
    y = BODY_TOP + 0.35
    for line in data["situation"]:
        h = text_height(line, left_w - 0.9, 16)
        add_box(s, MX + 0.35, y + 0.09, 0.12, 0.12, TEAL, MSO_SHAPE.OVAL)
        add_text(s, MX + 0.65, y, left_w - 1.0, h, line, 16)
        y += h + 0.28
    # 오른쪽: 분석 대상
    rx = MX + left_w + 0.4
    rw = SW - MX - rx
    add_text(s, rx, BODY_TOP, rw, 0.35, "분석 대상 보고서", 14, MUTED, bold=True)
    y = BODY_TOP + 0.5
    step = min(0.7, 4.3 / len(reports))
    for r in reports:
        add_box(s, rx, y, rw, step - 0.12, CARD)
        add_text(s, rx + 0.2, y, rw - 0.4, step - 0.12, r["filename"], 13, INK,
                 anchor=MSO_ANCHOR.MIDDLE)
        y += step


def build_problems(prs, data):
    n = len(data["problems"])
    for idx, p in enumerate(data["problems"], 1):
        cards = []
        cw = BODY_W - 0.6
        for e in p["evidence"]:
            h = 0.32 + text_height(e["quote"], cw, 14) + 0.3
            cards.append((h, e))
        note = re.sub(r"^\s*확인\s*필요\s*[:：]?\s*", "", (p.get("note") or "").strip())
        note_h = (text_height("확인 필요: " + note, BODY_W - 0.6, 13) + 0.3) if note else 0
        top = BODY_TOP + 0.95   # 제목 줄 + 부서 태그 줄
        avail = BODY_BOTTOM - top - (note_h + 0.15 if note else 0)
        pages = flow_pages(cards, avail, 0.15) or [[]]
        for pi, page in enumerate(pages):
            suffix = f" (계속 {pi + 1}/{len(pages)})" if len(pages) > 1 else ""
            s = new_slide(prs, "2", "확인된 문제", f"문제 {idx}/{n}")
            add_text(s, MX, BODY_TOP - 0.05, BODY_W, 0.5, p["title"] + suffix, 20, NAVY, bold=True)
            # 부서 태그
            x = MX
            for d in p["departments"]:
                w = 0.35 + len(d) * 0.19
                add_box(s, x, BODY_TOP + 0.55, w, 0.32, TEAL, radius=0.16)
                add_text(s, x, BODY_TOP + 0.55, w, 0.32, d, 12, WHITE, bold=True,
                         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                x += w + 0.12
            y = top
            for h, e in page:
                add_box(s, MX, y, BODY_W, h, CARD)
                add_text(s, MX + 0.3, y + 0.12, cw, 0.25, "근거: " + e["report"], 11, TEAL, bold=True)
                add_text(s, MX + 0.3, y + 0.42, cw, h - 0.5, e["quote"], 14)
                y += h + 0.15
            if note and pi == len(pages) - 1:
                add_box(s, MX, y, BODY_W, note_h, AMBER_TINT)
                add_text(s, MX + 0.3, y, BODY_W - 0.6, note_h, "확인 필요: " + note, 13, INK,
                         anchor=MSO_ANCHOR.MIDDLE)


def build_solutions(prs, data):
    n = len(data["solutions"])
    tag_w = 0.75
    for idx, sol in enumerate(data["solutions"], 1):
        tw = BODY_W - tag_w - 0.5
        rows = []
        for it in sol["items"]:
            body = it["text"]
            h = text_height(body, tw, 15) + (0.25 if it["tag"] == "원문" and it["source"] else 0) + 0.25
            rows.append((max(h, 0.6), it))
        avail = BODY_BOTTOM - (BODY_TOP + 0.6)
        pages = flow_pages(rows, avail, 0.12) or [[]]
        for pi, page in enumerate(pages):
            suffix = f" (계속 {pi + 1}/{len(pages)})" if len(pages) > 1 else ""
            s = new_slide(prs, "3", "해결 방안 제안", f"문제 {idx}/{n}")
            add_text(s, MX, BODY_TOP - 0.05, BODY_W, 0.5, sol["problem_title"] + suffix, 20, NAVY, bold=True)
            y = BODY_TOP + 0.6
            for h, it in page:
                is_prop = it["tag"] == "제안"
                add_box(s, MX, y, BODY_W, h, AMBER_TINT if is_prop else CARD)
                add_box(s, MX + 0.2, y + 0.15, tag_w, 0.32, AMBER if is_prop else GRAY_TAG, radius=0.16)
                add_text(s, MX + 0.2, y + 0.15, tag_w, 0.32, "[" + it["tag"] + "]", 12, WHITE, bold=True,
                         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                tx = MX + 0.2 + tag_w + 0.25
                add_text(s, tx, y + 0.12, tw, h - 0.2, it["text"], 15)
                if it["tag"] == "원문" and it["source"]:
                    add_text(s, tx, y + h - 0.32, tw, 0.22, "출처: " + it["source"], 10, MUTED)
                y += h + 0.12
            add_footer(s, "[제안] = 원문에 없는 새로운 제안   [원문] = 보고서에 이미 있는 조치")


def build_effects(prs, data):
    s = new_slide(prs, "4", "기대 효과")
    effects = data["effects"]
    cols = 2 if len(effects) > 3 else 1
    cw = (BODY_W - 0.3 * (cols - 1)) / cols
    rows = math.ceil(len(effects) / cols)
    ch = min(1.3, (4.7 - 0.25 * (rows - 1)) / rows)
    for i, e in enumerate(effects):
        c, r = i % cols, i // cols
        x = MX + c * (cw + 0.3)
        y = BODY_TOP + r * (ch + 0.25)
        add_box(s, x, y, cw, ch, CARD)
        add_box(s, x + 0.25, y + ch / 2 - 0.25, 0.5, 0.5, TEAL, MSO_SHAPE.OVAL)
        add_text(s, x + 0.25, y + ch / 2 - 0.25, 0.5, 0.5, str(i + 1), 16, WHITE, bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, x + 1.0, y + 0.1, cw - 1.25, ch - 0.2, e, 15, anchor=MSO_ANCHOR.MIDDLE)
    add_footer(s, "효과 수치는 보고서 원문에 없어 제시하지 않았습니다. 정성적 기대 효과입니다.")


def build_needs(prs, data):
    sched = data["known_schedule"]
    confirm = data["to_confirm"]
    lw = 6.0
    rw = BODY_W - lw - 0.4

    def chunks(items, heights, width, avail):
        return flow_pages(list(zip(heights, items)), avail, 0.1) or [[]]

    sh = [text_height(k["text"], lw - 0.6, 14) + 0.5 for k in sched]
    ch = [text_height(c, rw - 0.9, 14) + 0.3 for c in confirm]
    avail = BODY_BOTTOM - (BODY_TOP + 0.5)
    spages = chunks(sched, sh, lw, avail)
    cpages = chunks(confirm, ch, rw, avail)
    total = max(len(spages), len(cpages))
    for pi in range(total):
        suffix = f"{pi + 1}/{total}" if total > 1 else ""
        s = new_slide(prs, "5", "필요한 것과 일정", suffix)
        add_text(s, MX, BODY_TOP - 0.05, lw, 0.4, "원문에 있는 일정", 16, NAVY, bold=True)
        y = BODY_TOP + 0.5
        for h, k in (spages[pi] if pi < len(spages) else []):
            add_box(s, MX, y, lw, h, CARD)
            add_text(s, MX + 0.3, y + 0.12, lw - 0.6, h - 0.4, k["text"], 14)
            add_text(s, MX + 0.3, y + h - 0.3, lw - 0.6, 0.22, "출처: " + k["source"], 10, MUTED)
            y += h + 0.1
        rx = MX + lw + 0.4
        add_text(s, rx, BODY_TOP - 0.05, rw, 0.4, "확인 필요 사항", 16, NAVY, bold=True)
        y = BODY_TOP + 0.5
        for h, c in (cpages[pi] if pi < len(cpages) else []):
            add_box(s, rx, y, rw, h, AMBER_TINT)
            add_box(s, rx + 0.25, y + h / 2 - 0.09, 0.18, 0.18, AMBER, MSO_SHAPE.OVAL)
            add_text(s, rx + 0.65, y, rw - 0.9, h, c, 14, anchor=MSO_ANCHOR.MIDDLE)
            y += h + 0.1
        add_footer(s, "담당자·예산·신규 일정은 원문에 없어 정하지 않았습니다. 위 확인 필요 사항을 먼저 확정해야 합니다.")


def build_pptx(data, reports, path):
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    today = datetime.now().strftime("%Y-%m-%d")
    build_title_slide(prs, reports, today)
    build_situation(prs, data, reports)
    build_problems(prs, data)
    build_solutions(prs, data)
    build_effects(prs, data)
    build_needs(prs, data)
    prs.save(path)
    return len(prs.slides)


def make_output_path(folder: str) -> str:
    today = datetime.now().strftime("%Y-%m-%d")
    path = os.path.join(folder, f"공통문제_해결제안서_{today}.pptx")
    idx = 1
    while os.path.exists(path):
        path = os.path.join(folder, f"공통문제_해결제안서_{today}_{idx}.pptx")
        idx += 1
    return path


def main():
    folder = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__))

    try:
        api_key = load_api_key()
        reports = load_reports(folder)
        data = generate_proposal(api_key, reports)
        warnings = check_proposal(data, reports)
        output_path = make_output_path(folder)
        slide_count = build_pptx(data, reports, output_path)
    except (RuntimeError, APIError, APIStatusError) as e:
        print(f"오류가 발생하여 결과 파일을 만들지 않았습니다: {e}")
        return
    except Exception as e:
        print(f"예상하지 못한 오류가 발생하여 결과 파일을 만들지 않았습니다: {e}")
        return

    print(f"읽은 보고서 수: {len(reports)}")
    for r in reports:
        print(f" - {r['filename']}")
    print(f"결과 저장 위치: {output_path} (슬라이드 {slide_count}장)")
    if warnings:
        print("\n[점검 경고] 아래 항목은 사람이 직접 확인해주세요.")
        for w in warnings:
            print(f" ! {w}")
    else:
        print("\n[점검] 문제별 근거·인용 일치, 보고서 인용, '제안' 표시 확인 완료.")


if __name__ == "__main__":
    main()
